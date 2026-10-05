"""Strip DONE / IN PROGRESS status prefixes from the team report."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document

PATH = Path(r"c:\Users\avwhitehorn\Documents\GitHub\BessForge\docs\10.5.2026 Team Report.docx")
TMP = PATH.with_name("10.5.2026 Team Report.cleaned.docx")

PREFIX = re.compile(
    r"^(DONE(?:\s*/\s*ongoing polish)?|IN PROGRESS\s*/\s*watch)\s*[—–-]\s*",
    re.IGNORECASE,
)


def set_runs(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def main() -> None:
    doc = Document(str(PATH))
    changed = 0
    for p in doc.paragraphs:
        original = p.text
        cleaned = PREFIX.sub("", original)
        if cleaned != original:
            set_runs(p, cleaned)
            changed += 1
    doc.save(str(TMP))
    print(f"Wrote cleaned copy with {changed} paragraphs updated: {TMP}")
    try:
        shutil.copyfile(TMP, PATH)
        TMP.unlink(missing_ok=True)
        print(f"Replaced original: {PATH}")
    except PermissionError:
        print(
            "ORIGINAL LOCKED — close the Word doc, then replace it with:\n"
            f"  {TMP}"
        )


if __name__ == "__main__":
    main()
